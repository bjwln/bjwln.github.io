import sys, io, os, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH

dates = [
    "9 月16 日","9 月23 日","10 月14 日","10 月21 日","10 月28 日",
    "11 月4 日","11 月11 日","11 月18 日","11 月25 日","12 月2 日",
    "12 月9 日","12 月16 日","12 月23 日","12 月30 日","1 月6 日","1 月13 日"
]

# 两组各5人，隔周轮换
# A组: 3研三 + 2研一 (赵春雨在A组→第1周)
group_a = ["黄坤","王炎吉","叶文斌","许锋敏","赵春雨"]
# B组: 4研二 + 1研一
group_b = ["何东","丁思涵","卢惠敏","黄星宇","陈天e"]

grades = {"黄坤":"研三","王炎吉":"研三","叶文斌":"研三",
          "何东":"研二","丁思涵":"研二","卢惠敏":"研二","黄星宇":"研二",
          "许锋敏":"研一","赵春雨":"研一","陈天翊":"研一"}

# A组奇数周(1,3,5,...15), B组偶数周(2,4,...16)
schedule = []
for w in range(16):
    if w % 2 == 0:  # 奇数周 (0-indexed: 0,2,4...)
        people = group_a
    else:
        people = group_b
    schedule.append((dates[w], people[:]))

# 统计
print("=== 排班表 ===")
for w in range(16):
    ps = schedule[w][1]
    date = schedule[w][0]
    grade_info = ", ".join(f"{p}({grades[p]})" for p in ps)
    print(f"  W{w+1:>2} {date:<10} {grade_info}")

print("\n=== 统计 ===")
for p in group_a + group_b:
    weeks_list = [w+1 for w in range(16) if p in schedule[w][1]]
    gaps = [weeks_list[i+1]-weeks_list[i] for i in range(len(weeks_list)-1)]
    print(f"  {p}({grades[p]}): {len(weeks_list)}次, 间隔{gaps}")

# 写入docx
doc = Document(r"C:/Users/lenovo/Desktop/新建 DOC 文档.docx")
table = doc.tables[0]

FONT_CN = "SimSun"
FONT_EN = "Times New Roman"

def set_run_font(run):
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), FONT_EN)
    rFonts.set(qn("w:hAnsi"), FONT_EN)
    rFonts.set(qn("w:eastAsia"), FONT_CN)
    rFonts.set(qn("w:cs"), FONT_EN)

def clear_cell(cell):
    paras = cell.paragraphs
    for p in paras[1:]:
        p._element.getparent().remove(p._element)
    p = paras[0]
    for run in list(p.runs):
        run._element.getparent().remove(run._element)
    return p

def fill_cell(cell, text, bold_digits=False):
    p = clear_cell(cell)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    segments = re.findall(r'[\d]+|[\u4e00-\u9fff、，（）]+|.', text)
    for seg in segments:
        run = p.add_run(seg)
        run.font.size = Pt(12)
        run.font.bold = (bold_digits and seg.isdigit())
        set_run_font(run)

# 表头
fill_cell(table.rows[0].cells[0], "")
fill_cell(table.rows[0].cells[1], "小组组会（每周三下午14:00开始）")

# 数据行
for i, (date, people) in enumerate(schedule):
    row = table.rows[i + 1]
    fill_cell(row.cells[0], date, bold_digits=True)
    fill_cell(row.cells[1], "、".join(people), bold_digits=False)

doc.save(r"C:/Users/lenovo/Desktop/新建 DOC 文档.docx")
print("\n已写入 docx!")
